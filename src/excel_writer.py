"""生成人工核查用 Excel：Items / ECR Tree / Edges / Metadata 四个 sheet。

- Items：全部实体平铺，列 = [ID, Kind, Level] + 字段表别名（剥 ALM_ 前缀），
  日期列（别名以 "Date" 结尾）写成真正的日期单元格并套用格式。
- ECR Tree：透视树。按 ECR 层级缩进 + Excel 大纲分组（原生 +/- 折叠展开），
  对 pivot_fields 指定的数值列做子树 sum（Subtotal 列）。
- Edges：显式关系边。
- Metadata：导出时间与各类计数。
"""
from __future__ import annotations
import os
from datetime import date, datetime
from typing import Any, Iterable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import Outline

import config
from data_utils import build_alias, parse_date

EXCEL_DATE_FORMAT = "yyyy-mm-dd"

TREE_SHEET = "ECR Tree"
TREE_BASE_COLUMNS = ["ID", "Kind", "Summary", "State", "Subtree Items"]

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_sheet(ws, columns: list[str], rows: Iterable[dict[str, Any]], date_columns: set[str]) -> None:
    data = list(rows)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = BORDER
    for row in data:
        values = [parse_date(row.get(c, "")) if c in date_columns else row.get(c, "") for c in columns]
        ws.append(values)
        for idx, cell in enumerate(ws[ws.max_row], start=1):
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if columns[idx - 1] in date_columns and isinstance(cell.value, (date, datetime)):
                cell.number_format = EXCEL_DATE_FORMAT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, column in enumerate(columns, start=1):
        samples = [column] + [str(row.get(column, "")) for row in data[:200]]
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max(map(len, samples)) + 2, 12), 60)


def write_items_to_excel(output_path, fields, items, relationships, pivot_fields=None) -> None:
    """items: [(kind, rows, levels)]，rows 的键是展示别名；relationships: {relation: edges}。

    pivot_fields: 透视树要 sum 的原始字段名列表（来自 ECR_Config.xlsx 的 'Pivot'
    sheet）；None 时缺省 config.DEFAULT_PIVOT_FIELD。不在字段表里的项会被忽略。
    """
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)

    aliases = [build_alias(field) for field in fields]
    item_columns = ["ID", "Kind", "Level", config.ALIAS_PROCESS_AREA] + [
        alias for alias in aliases if alias != "ID"]
    item_rows: list[dict[str, Any]] = []
    for kind, rows, levels in items:
        level_map = levels or {}
        for row in rows:
            item_id = str(row.get("ID", "")).strip()
            level_value = level_map.get(int(item_id), "") if item_id.isdigit() else ""
            item_row: dict[str, Any] = {
                "ID": item_id, "Kind": kind, "Level": level_value,
                config.ALIAS_PROCESS_AREA: row.get(config.ALIAS_PROCESS_AREA, ""),
            }
            for alias in aliases:
                if alias != "ID":
                    item_row[alias] = row.get(alias, "")
            item_rows.append(item_row)

    edge_columns = ["Parent ID", "Child ID", "Relation"]
    edge_rows = [
        {"Parent ID": parent_id, "Child ID": child_id, "Relation": relation}
        for relation, edges in relationships.items()
        for parent_id, child_id in edges
    ]

    wb = openpyxl.Workbook()
    ws_items = wb.active
    ws_items.title = "Items"
    _write_sheet(ws_items, item_columns, item_rows, {c for c in item_columns if c.endswith("Date")})

    _write_tree_sheet(wb, fields, items, relationships, pivot_fields)

    ws_edges = wb.create_sheet("Edges")
    _write_sheet(ws_edges, edge_columns, edge_rows, set())

    meta = wb.create_sheet("Metadata")
    meta.append(["Export Time", datetime.now()])
    meta["B1"].number_format = "yyyy-mm-dd hh:mm:ss"
    meta.append(["Item kind", "Count"])
    for kind, rows, _ in items:
        meta.append([kind, len(rows)])
    meta.append(["Relation", "Count"])
    for relation, edges in relationships.items():
        meta.append([relation, len(edges)])

    wb.save(output_path)


# ---------------------------------------------------------------------------
# ECR Tree（透视树）
# ---------------------------------------------------------------------------

def _to_number(value: Any) -> float | None:
    """尽力把单元格值转成 float（容忍千分位逗号）；失败返回 None。"""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _int_id(row: dict[str, Any]) -> int | None:
    text = str(row.get("ID", "")).strip()
    return int(text) if text.isdigit() else None


def _children_index(relationships: dict[str, list[tuple[int, int]]]) -> dict[str, dict[int, list[int]]]:
    """{relation: {parent_id: [child_id, ...]}}，去重、保持边顺序。"""
    index: dict[str, dict[int, list[int]]] = {}
    for relation, edges in relationships.items():
        mapping: dict[int, list[int]] = {}
        for parent_id, child_id in edges:
            children = mapping.setdefault(parent_id, [])
            if child_id not in children:
                children.append(child_id)
        index[relation] = mapping
    return index


def _write_tree_sheet(wb, fields, items, relationships, pivot_fields) -> None:
    """按 ECR 逐棵深度优先输出层级行，配 Excel 大纲分组实现 +/- 折叠。

    口径（详见 Handoff/data-pipeline.md §5.2）：
    - Subtotal = 自身 + 全部后代（每个 ECR 子树内节点去重，多父级只挂在首条路径下）；
    - 同一节点被不同 ECR 引用时各出现一次，TOTAL = 各 ECR Subtotal 之和，
      因此共享节点在 TOTAL 中会被计多次（列内一致性优先）。
    """
    alias_set = {build_alias(field) for field in fields}
    requested = list(pivot_fields) if pivot_fields is not None else [config.DEFAULT_PIVOT_FIELD]
    metrics = [build_alias(field) for field in requested if build_alias(field) in alias_set]

    rows_by_id: dict[int, tuple[str, dict[str, Any]]] = {}
    order_pos: dict[int, int] = {}
    global_pos = 0
    for kind, rows, _ in items:
        for row in rows:
            item_id = _int_id(row)
            if item_id is not None:
                rows_by_id[item_id] = (kind, row)
                order_pos.setdefault(item_id, global_pos)
                global_pos += 1

    children = _children_index(relationships)
    wi_children = children.get(config.REL_WORK_ITEMS, {})
    action_children = children.get(config.REL_ACTIONS, {})
    build_children = children.get(config.REL_WORK_ITEM_FOR, {})

    def child_ids_of(item_id: int, kind: str) -> list[int]:
        ids: list[int] = []
        if kind in (config.KIND_ECR, config.KIND_WORK_ITEM):
            ids.extend(wi_children.get(item_id, []))
        if kind == config.KIND_ECR:
            ids.extend(action_children.get(item_id, []))
            ids.extend(build_children.get(item_id, []))
        # edges 集合是排序过的，不反映导出顺序；子节点按 Items 发现顺序输出
        return sorted(set(ids), key=lambda cid: order_pos.get(cid, len(order_pos)))

    def build_node(item_id: int, depth: int, visited: set[int]) -> dict[str, Any]:
        kind, row = rows_by_id[item_id]
        visited.add(item_id)
        own_num: list[float] = []
        own_display: list[Any] = []
        for metric in metrics:
            raw = row.get(metric, "")
            number = _to_number(raw)
            own_num.append(0.0 if number is None else number)
            own_display.append(number if number is not None else raw)
        nodes = []
        for child_id in child_ids_of(item_id, kind):
            if child_id in visited or child_id not in rows_by_id:
                continue
            nodes.append(build_node(child_id, depth + 1, visited))
        subtotal = list(own_num)
        count = 1
        for child in nodes:
            subtotal = [a + b for a, b in zip(subtotal, child["subtotal"])]
            count += child["count"]
        return {"id": item_id, "kind": kind, "row": row, "depth": depth,
                "own_display": own_display, "subtotal": subtotal,
                "count": count, "children": nodes}

    ecr_nodes = []
    for kind, rows, _ in items:
        if kind != config.KIND_ECR:
            continue
        for row in rows:
            item_id = _int_id(row)
            if item_id is not None:
                ecr_nodes.append(build_node(item_id, 0, set()))

    columns = TREE_BASE_COLUMNS + [c for metric in metrics for c in (metric, f"{metric} Subtotal")]
    ws = wb.create_sheet(TREE_SHEET)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = BORDER

    total_count = sum(node["count"] for node in ecr_nodes)
    total_values: list[Any] = ["", "", "TOTAL (all ECRs)", "", total_count]
    for i in range(len(metrics)):
        total_values.extend(["", round(sum(node["subtotal"][i] for node in ecr_nodes), 6)])
    ws.append(total_values)
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.border = BORDER

    def emit(node: dict[str, Any]) -> None:
        values: list[Any] = [node["id"], node["kind"], node["row"].get("Summary", ""),
                             node["row"].get("State", ""), node["count"]]
        for i in range(len(metrics)):
            values.append(node["own_display"][i])
            values.append(round(node["subtotal"][i], 6))
        ws.append(values)
        excel_row = ws.max_row
        if node["depth"] > 0:
            ws.row_dimensions[excel_row].outlineLevel = node["depth"]
        for cell in ws[excel_row]:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        # Summary 列按深度缩进，折叠展开时层级一目了然
        ws.cell(excel_row, 3).alignment = Alignment(vertical="top", wrap_text=True, indent=node["depth"])
        if node["kind"] == config.KIND_ECR:
            for cell in ws[excel_row]:
                cell.font = Font(bold=True)
        for child in node["children"]:
            emit(child)

    for node in ecr_nodes:
        emit(node)

    # summaryBelow=False：组的汇总行（父节点）在上方，+/− 按钮出现在父行
    ws.sheet_properties.outlinePr = Outline(summaryBelow=False)
    ws.freeze_panes = "A3"
    for idx, column in enumerate(columns, start=1):
        samples = [column]
        for row_vals in ws.iter_rows(min_row=2, min_col=idx, max_col=idx, values_only=True):
            samples.append("" if row_vals[0] is None else str(row_vals[0]))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max(map(len, samples)) + 2, 12), 60)
