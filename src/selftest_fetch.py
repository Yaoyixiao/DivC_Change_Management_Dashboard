"""离线冒烟测试：用假数据替代 im exportissues，全链路验证 fetcher。

用法
----
    python src/selftest_fetch.py

不依赖 PTC 客户端与网络：替换 fetcher 的导出/读取入口（假导出永远写 xlsx
内容），把配置与输出路径指到临时目录，跑完 fetcher.main() 后断言
SQLite / Excel / manifest / schema 的内容，最后自动清理。

真实取数请运行 fetcher.py；本脚本只验证流水线本身，不验证 PTC 侧行为。

假数据覆盖的场景
----------------
- 多个 ECR 共享同一个 Work Item（只导出一次，两条边都保留）
- Work Item 三级递归（L0/L1/L2），最深层不再下钻
- 关联目标不存在（边被丢弃）、根 ID 不存在（计数告警）
- 最深层回指上层节点（不产生环）
- ISO 字符串日期与 Excel 序列数日期两种写法
"""
from __future__ import annotations

import io
import json
import re
import sqlite3
import sys
import tempfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl
from openpyxl.utils.datetime import from_excel

import config
import fetcher
from data_utils import build_alias
from parser import read_xlsx_rows

ROOT_IDS = [101, 102, 103, 109]  # 109 在 PTC 不存在

FIELDS = [
    "ID", "Type", "Summary", "State", "Created Date", "ALM_Owners",
    "ALM_Planned Start Date", "ALM_Target Date", "ALM_Maturity Level",
    "ALM_Planned Effort", "ALM_Work Items", "ALM_Actions", "ALM_Work Item For",
]

# 模拟 im exportissues 的原始导出：键为原始字段名
FAKE_ITEMS: dict[int, dict] = {
    101: {"ID": 101, "Type": "ECR", "Summary": "ECR 101", "State": "ALM_Defined",
          "Created Date": "2026-01-15", "ALM_Target Date": "2026-06-30",
          "ALM_Work Items": "201, 202", "ALM_Actions": "301", "ALM_Work Item For": "401"},
    102: {"ID": 102, "Type": "ECR", "Summary": "ECR 102", "State": "ALM_Initiated",
          "Created Date": "2026-02-01",
          "ALM_Work Items": "202, 203, 999",
          "ALM_Actions": "302, 399",
          "ALM_Work Item For": "402"},
    103: {"ID": 103, "Type": "ECR", "Summary": "ECR 103", "State": "ALM_Closed",
          "Created Date": "2026-02-10"},
    201: {"ID": 201, "Type": "Work Item", "Summary": "WI 201", "State": "ALM_Started",
          "ALM_Planned Start Date": "2026-03-15", "ALM_Work Items": "210"},
    202: {"ID": 202, "Type": "Work Item", "Summary": "WI 202", "State": "ALM_Defined",
          "ALM_Planned Start Date": 46000, "ALM_Work Items": "210"},
    210: {"ID": 210, "Type": "Work Item", "Summary": "WI 210", "State": "ALM_Planned",
          "ALM_Work Items": "220"},
    220: {"ID": 220, "Type": "Work Item", "Summary": "WI 220", "State": "ALM_Initiated",
          "ALM_Work Items": "201, 230"},  # 已到 L2：不下钻；201 是回指
    230: {"ID": 230, "Type": "Work Item", "Summary": "WI 230", "State": "ALM_Initiated"},
    301: {"ID": 301, "Type": "Action", "Summary": "Action 301", "State": "ALM_Checked"},
    302: {"ID": 302, "Type": "Action", "Summary": "Action 302", "State": "ALM_Closed"},
    401: {"ID": 401, "Type": "Build", "Summary": "Build 401", "ALM_Target Date": "2026-08-01",
          "ALM_Maturity Level": "CAT 3"},
    402: {"ID": 402, "Type": "Build", "Summary": "Build 402", "ALM_Target Date": "2026-09-15",
          "ALM_Maturity Level": "CAT 4"},
}

EXPECTED_WI_EDGES = {
    (101, 201), (101, 202), (102, 202),  # 202 被两个 ECR 共享
    (201, 210), (202, 210),              # 210 被两个 WI 共享（同一 hop）
    (210, 220),
}
EXPECTED_ACTION_EDGES = {(101, 301), (102, 302)}   # (102, 399) 目标不存在，丢弃
EXPECTED_BUILD_EDGES = {(101, 401), (102, 402)}

EXPECTED_COUNTS = {"ecr": 3, "work_item": 4, "action": 2, "build": 2}
EXPECTED_RELATIONSHIPS = {
    config.REL_WORK_ITEMS: len(EXPECTED_WI_EDGES),
    config.REL_ACTIONS: len(EXPECTED_ACTION_EDGES),
    config.REL_WORK_ITEM_FOR: len(EXPECTED_BUILD_EDGES),
}


def fake_export_issues(query_definition, fields, output_file) -> None:
    """替代 cli_client.export_issues：按查询里的 ID 写 xlsx 内容（后缀名不变）。"""
    ids = [int(value) for value in re.findall(r"field\[ID\]=(\d+)", query_definition)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(fields)
    for item_id in ids:
        item = FAKE_ITEMS.get(item_id)
        if item is None:
            continue  # 模拟 PTC 静默跳过不存在的 ID
        ws.append([item.get(field, "") for field in fields])
    wb.save(output_file)


def write_fake_config(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ECR"
    ws.append(["ID"])
    for root_id in ROOT_IDS:
        ws.append([root_id])
    ws_fields = wb.create_sheet("Data_Field")
    ws_fields.append(["Field"])
    for field in FIELDS:
        ws_fields.append([field])
    wb.save(path)


def fake_read_exported_excel(path):
    """假导出写入的是 xlsx 内容但后缀仍为 .xls（openpyxl 按后缀拒载），
    这里把字节流经 BytesIO 交给 read_xlsx_rows 按内容读取。"""
    return read_xlsx_rows(io.BytesIO(Path(path).read_bytes()))


def run_fetcher_with_fakes(tmp_path: Path) -> None:
    config_path = tmp_path / "ECR_Config.xlsx"
    write_fake_config(config_path)

    path_patches = {
        "ECR_CONFIG_PATH": str(config_path),
        "OUTPUT_PATH": str(tmp_path / "data.xlsx"),
        "DATABASE_PATH": str(tmp_path / "dashboard.db"),
        "MANIFEST_PATH": str(tmp_path / "dashboard_manifest.json"),
        "SCHEMA_PATH": str(tmp_path / "dashboard_schema.json"),
        "TEMP_DIR": str(tmp_path / ".tmp_exportissues"),
    }
    original_config = {key: getattr(config, key) for key in path_patches}
    original_export_issues = fetcher.export_issues
    original_read = fetcher.read_exported_excel
    try:
        for key, value in path_patches.items():
            setattr(config, key, value)
        fetcher.export_issues = fake_export_issues
        fetcher.read_exported_excel = fake_read_exported_excel
        fetcher.main()
    finally:
        for key, value in original_config.items():
            setattr(config, key, value)
        fetcher.export_issues = original_export_issues
        fetcher.read_exported_excel = original_read


def assert_database(tmp_path: Path) -> None:
    con = sqlite3.connect(tmp_path / "dashboard.db")
    try:
        cur = con.cursor()

        counts = dict(cur.execute("SELECT kind, COUNT(*) FROM items GROUP BY kind").fetchall())
        assert counts == EXPECTED_COUNTS, f"kind 计数不符: {counts}"

        ghosts = cur.execute(
            "SELECT COUNT(*) FROM items WHERE id IN (109, 203, 999, 230, 399)"
        ).fetchone()[0]
        assert ghosts == 0, f"不存在的 ID 落库了: {ghosts} 条"

        levels = dict(cur.execute(
            "SELECT id, level FROM items WHERE kind = 'work_item'").fetchall())
        assert levels == {201: 0, 202: 0, 210: 1, 220: 2}, f"level 不符: {levels}"
        ecr_with_level = cur.execute(
            "SELECT COUNT(*) FROM items WHERE kind = 'ecr' AND level IS NOT NULL").fetchone()[0]
        assert ecr_with_level == 0, "非 work_item 不应有 level"

        def edges(relation):
            return set(cur.execute(
                "SELECT parent_id, child_id FROM edges WHERE relation = ?", (relation,)).fetchall())

        assert edges(config.REL_WORK_ITEMS) == EXPECTED_WI_EDGES, edges(config.REL_WORK_ITEMS)
        assert edges(config.REL_ACTIONS) == EXPECTED_ACTION_EDGES, edges(config.REL_ACTIONS)
        assert edges(config.REL_WORK_ITEM_FOR) == EXPECTED_BUILD_EDGES, edges(config.REL_WORK_ITEM_FOR)

        meta = dict(cur.execute("SELECT key, value FROM metadata").fetchall())
        assert json.loads(meta["root_item_ids"]) == ROOT_IDS, meta.get("root_item_ids")
        assert meta["config_file"] == "ECR_Config.xlsx"
        assert "generated_at" in meta and "schema_version" in meta

        dates = dict(cur.execute(
            "SELECT id, planned_start_date FROM items WHERE kind = 'work_item'").fetchall())
        assert dates[201] == "2026-03-15", dates
        assert dates[202] == from_excel(46000).date().isoformat(), dates  # Excel 序列数 -> ISO
        created = dict(cur.execute(
            "SELECT id, created_date FROM items WHERE kind = 'ecr'").fetchall())
        assert created[101] == "2026-01-15", created

        columns = [row[1] for row in cur.execute("PRAGMA table_info(items)").fetchall()]
        assert "planned_start_date" in columns and "maturity_level" in columns, columns
        assert "kind" in columns and "level" in columns, columns
    finally:
        con.close()


def assert_excel(tmp_path: Path) -> None:
    wb = openpyxl.load_workbook(tmp_path / "data.xlsx")
    assert wb.sheetnames == ["Items", "Edges", "Metadata"], wb.sheetnames

    ws = wb["Items"]
    headers = [cell.value for cell in ws[1]]
    expected_headers = ["ID", "Kind", "Level"] + [
        build_alias(field) for field in FIELDS if build_alias(field) != "ID"]
    assert headers == expected_headers, headers
    data_rows = ws.max_row - 1
    assert data_rows == sum(EXPECTED_COUNTS.values()), data_rows

    by_id = {int(row[0]): row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert by_id[101][1] == "ecr" and by_id[101][2] in ("", None), by_id[101][:3]
    assert by_id[201][1] == "work_item" and by_id[201][2] == 0, by_id[201][:3]
    assert by_id[220][2] == 2, by_id[220][:3]
    assert isinstance(by_id[201][headers.index("Planned Start Date")], datetime), "日期列应为日期单元格"

    ws_edges = wb["Edges"]
    assert ws_edges.max_row - 1 == sum(EXPECTED_RELATIONSHIPS.values()), ws_edges.max_row

    ws_meta = wb["Metadata"]
    meta_values = {(row[0], row[1]) for row in ws_meta.iter_rows(values_only=True)}
    assert ("work_item", 4) in meta_values and ("work_items", 6) in meta_values, meta_values


def assert_json(tmp_path: Path) -> None:
    manifest = json.loads((tmp_path / "dashboard_manifest.json").read_text(encoding="utf-8"))
    assert manifest["status"] == "ready"
    assert manifest["config_file"] == "ECR_Config.xlsx"
    assert manifest["counts"] == EXPECTED_COUNTS, manifest["counts"]
    assert manifest["relationship_counts"] == EXPECTED_RELATIONSHIPS, manifest["relationship_counts"]

    schema = json.loads((tmp_path / "dashboard_schema.json").read_text(encoding="utf-8"))
    item_columns = schema["tables"]["items"]["columns"]
    assert item_columns["id"] == "integer" and item_columns["level"] == "integer"
    assert item_columns["planned_start_date"] == "date"
    assert item_columns["created_date"] == "date"
    assert item_columns["summary"] == "text"
    assert schema["kinds"] == list(config.KINDS)
    assert schema["relations"] == [config.REL_WORK_ITEMS, config.REL_ACTIONS, config.REL_WORK_ITEM_FOR]


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="ecr_selftest_") as tmp:
        tmp_path = Path(tmp)
        run_fetcher_with_fakes(tmp_path)
        assert_database(tmp_path)
        assert_excel(tmp_path)
        assert_json(tmp_path)
        assert not (tmp_path / ".tmp_exportissues").exists(), "临时导出目录未清理"
    print("selftest_fetch: PASS (db / xlsx / manifest / schema)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
